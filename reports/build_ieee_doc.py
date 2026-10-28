from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent
lines = (root / 'IEEE_Research_Report_Draft.md').read_text(encoding='utf-8').splitlines()
d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.65)
normal = d.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.0
for name in ('Title', 'Heading 1', 'Heading 2'):
    s = d.styles[name]
    s.font.name = 'Times New Roman'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(6)
    s.paragraph_format.space_after = Pt(3)
d.styles['Title'].font.size = Pt(20)
d.styles['Heading 1'].font.size = Pt(10)
d.styles['Heading 2'].font.size = Pt(10)

def rich(p, text):
    for i, part in enumerate(re.split(r'(\*\*.*?\*\*|\*[^*]+\*)', text)):
        if part.startswith('**') and part.endswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*'):
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)

i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line or line == '---':
        i += 1; continue
    if line.startswith('# '):
        p = d.add_paragraph(line[2:], 'Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ppr = p._p.get_or_add_pPr()
        borders = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom'); bottom.set(qn('w:val'), 'nil'); borders.append(bottom); ppr.append(borders)
        for run in p.runs:
            run.font.name = 'Times New Roman'; run.font.size = Pt(20); run.font.color.rgb = RGBColor(0, 0, 0)
    elif line.startswith('**Author(s):') or line.startswith('**Affiliation:') or line.startswith('**Email:') or line.startswith('**Status:'):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rich(p, line)
    elif line.startswith('**Abstract'):
        p = d.add_paragraph(); rich(p, line); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif line.startswith('**Index Terms'):
        p = d.add_paragraph(); rich(p, line)
        s = d.add_section(WD_SECTION.CONTINUOUS)
        cols = s._sectPr.xpath('./w:cols')[0]; cols.set(qn('w:num'), '2'); cols.set(qn('w:space'), '360')
    elif line.startswith('## '):
        p = d.add_paragraph(line[3:], 'Heading 1'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith('### '):
        p = d.add_paragraph(line[4:], 'Heading 2'); p.runs[0].italic = True
    elif line.startswith('| Model'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(set(x) <= set('-: ') for x in cells): rows.append(cells)
            i += 1
        rows[0] = ['Model', 'HAM acc.', 'HAM F1', 'ISIC acc.', 'ISIC F1']
        rows[-1][0] = "Authors' reference"
        t = d.add_table(rows=0, cols=5); t.autofit = True; t.style = 'Table Grid'
        for row in rows:
            c = t.add_row().cells
            for j, val in enumerate(row):
                c[j].text = val
                for p in c[j].paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for run in p.runs: run.font.size = Pt(7)
        continue
    else:
        p = d.add_paragraph(); rich(p, line)
        if not line.startswith('**TABLE') and not line.startswith('*Drafting'): p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    i += 1
d.save(root / 'IEEE_Research_Report_Draft.docx')
