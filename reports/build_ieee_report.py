from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path('reports/IEEE_Skin_Lesion_Classification_Report_Detailed.docx')
doc=Document(); sec=doc.sections[0]
for s in doc.sections:
    s.top_margin=Inches(.58); s.bottom_margin=Inches(.58); s.left_margin=Inches(.62); s.right_margin=Inches(.62)
normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(10)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
normal.paragraph_format.space_after=Pt(3); normal.paragraph_format.line_spacing=1.0
for n,z in [('Heading 1',11),('Heading 2',10)]:
    st=doc.styles[n]; st.font.name='Times New Roman'; st.font.size=Pt(z); st.font.bold=True; st.font.color.rgb=None
    st.paragraph_format.space_before=Pt(7); st.paragraph_format.space_after=Pt(3); st.paragraph_format.keep_with_next=True

def para(t, label=None, center=False):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if label:
        r=p.add_run(label); r.bold=True
    p.add_run(t); return p
def h(t,l=1): doc.add_heading(t,l)
def table(caption, headers, rows):
    if caption.startswith('TABLE'): para(caption,center=True)
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row): c.text=str(v)
    for index,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                if index < len(t.rows)-1: p.paragraph_format.keep_with_next=True
                for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(8)
    if not caption.startswith('TABLE'): para(caption,center=True)
    return t

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('A Comparative Study of Frozen Visual Representations for\nSeven-Class Dermoscopic Lesion Classification'); r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(18)
para('AI Project Research Team  |  Independent Research Prototype  |  October 2026',center=True)
para('This report documents the architecture, experimental results, and implementation status of a research prototype for seven-class dermoscopic lesion classification. The project compares frozen ImageNet EfficientNet-B0 features and frozen PanDerm vision-transformer representations with linear, multilayer perceptron (MLP), gradient-boosted tree, and one-versus-rest specialist classifiers. A separate federated-learning path implements Flower, FedAvg/FedProx options, and optional Opacus differential privacy. On the official ISIC 2018 Task 3 test cohort, PanDerm Base embeddings with an MLP achieved 80.56% accuracy and 70.30% macro-F1; a local reproduction of the PanDerm Large authors’ linear-probe recipe achieved 82.94% accuracy and 71.58% macro-F1. These are historical, repeatedly inspected benchmark results, not fresh blind validation. Improved EfficientNet features achieved 61.46% macro-F1. A later PanDerm adaptation used both former test cohorts for additional fitting, so those scores are fitting diagnostics only. The federated path was demonstrated only with synthetic data and did not establish lesion-classification utility. The prototype is not intended for clinical use.',label='Abstract—')
para('dermoscopy, skin lesion classification, transfer learning, vision transformer, EfficientNet, PanDerm, class imbalance, federated learning.',label='Index Terms—')
body=doc.add_section(WD_SECTION.CONTINUOUS); body.top_margin=Inches(.58); body.bottom_margin=Inches(.58); body.left_margin=Inches(.62); body.right_margin=Inches(.62)
cols=body._sectPr.find(qn('w:cols'))
if cols is None: cols=OxmlElement('w:cols'); body._sectPr.append(cols)
cols.set(qn('w:num'),'2'); cols.set(qn('w:space'),'360')

h('I. INTRODUCTION')
para('This project studies image classification over seven diagnoses: actinic keratoses (akiec), basal cell carcinoma (bcc), benign keratosis-like lesions (bkl), dermatofibroma (df), melanoma (mel), melanocytic nevus (nv), and vascular lesion (vasc). Its central question is how well fixed pretrained visual representations support these classes using smaller task-specific classifiers.')
para('The codebase grew from a federated CNN prototype into a local research workbench. Evidence spans independent experiment families rather than one matched leaderboard. The report keeps frozen-feature studies separate from the exploratory federated simulation and identifies the protocol alongside each result.')
para('The work provides a reproducible embedding and classification pipeline; comparisons among linear, nonlinear, tree, and specialist heads; a verified local PanDerm encoder and authors’ probe reference; and a Flower/Opacus prototype whose measured federated evidence is limited to synthetic smoke data.')
para('The project’s primary outcome is a research comparison across model families and operating metrics, not a clinical claim. In the frozen PanDerm Large comparison, the custom linear probe achieved 73.80% macro-F1 on the repeatedly inspected official ISIC cohort; the unchanged authors’ probe achieved 71.58%. These numbers describe a particular local protocol and cannot establish prospective superiority. The following sections detail how the values were produced and what they can and cannot support [5].')
h('II. RELATED WORK')
para('HAM10000 is a dermatoscopic image collection with seven diagnosis categories [1]. ISIC 2018 Task 3 provides a lesion-diagnosis evaluation setting [2]. PanDerm is a multimodal dermatology foundation model pretrained across clinical imaging modalities and institutions; this project uses released image encoders as fixed representation extractors [3]. EfficientNet is used as an ImageNet-pretrained convolutional feature extractor [4].')
h('III. SYSTEM ARCHITECTURE')
h('A. Image Classification Flow',2)
para('The inference path is: dermoscopic image → RGB decoding and deterministic preprocessing → frozen encoder → feature vector → classifier head → seven class scores and predicted label. The application loads saved model artifacts and displays predictions and results. Training and inference code paths are separate; artifacts retain parameters, configuration, metrics, and image-level predictions.')
table('Fig. 1. Centralized classification path. Federated training is a separate experimental branch.', ['Input','Preprocess','Frozen encoder','Embedding','Classifier','Output'], [['Image','Resize/crop\nnormalize','EfficientNet or\nPanDerm','1280 / 768 /\n1024 values','Linear / MLP /\nXGBoost','Seven scores\nand class']])
h('B. Representation Backbones',2)
para('EfficientNet-B0 uses ImageNet-pretrained weights in evaluation mode without gradient updates and emits 1,280-dimensional descriptors. In the improved multi-view run, a canonical crop descriptor is concatenated with the mean of eight square-symmetry full-image view descriptors, yielding 2,560 dimensions. Cached features support classifier trials without repeatedly running the encoder.')
para('PanDerm Base and Large emit 768- and 1,024-dimensional embeddings. The Large encoder has 303,326,208 parameters and remains frozen in local head comparisons. Preprocessing follows the pinned upstream evaluation transform: RGB, resize to 256 pixels, center crop to 224 pixels, and channel normalization. Encoder source and checkpoint tensors are hash-checked and loaded strictly. The measured system is encoder plus classifier, not a bare zero-shot classifier.')
para('The PanDerm input transform uses channel means (0.485, 0.456, 0.406) and standard deviations (0.228, 0.224, 0.225), including the upstream red-channel value of 0.228. The Base model is ViT-B/16 with 12 transformer blocks; Large uses 24 blocks. Features are computed in evaluation mode with no gradient updates. The implementation checks that encoder keys and tensor shapes match the expected model before loading the checkpoint. Each embedding can then be reused by multiple heads, avoiding a costly encoder pass for every classifier trial.')
h('C. Classifier Heads',2)
para('Linear probe: multinomial logistic regression. Custom runs use training-fitted scaling and class weighting; the PanDerm authors’ reference preserves the released raw-feature probe with C=71.68. MLP: LayerNorm, ReLU, dropout, and seven output logits. Representative shapes are 1280→256→64→7, 1024→256→64→7, and 1024→192→64→7, trained with AdamW and validation early stopping. XGBoost: shallow boosted trees over frozen features, with class weights and validation stopping; one PanDerm experiment used supervised six-dimensional LDA before the trees. Lesion specialists: seven independently selected one-versus-rest classifiers with fold-derived calibration and thresholds; their binary metrics differ from final seven-way argmax metrics.')
para('Scaling, projection, and model selection are fitted within training partitions or folds. Validation selects settings as declared for each experiment. Validation and test cohorts have nevertheless been examined repeatedly across project experiments.')
para('For a linear head, z = Wx + b is converted to seven class scores by the softmax function, pᵢ = exp(zᵢ)/Σⱼexp(zⱼ). A neural head inserts one or more learned nonlinear layers before the final logits. The selected Large 192→64 network contains 210,119 trainable head parameters and was chosen in a three-fold lesion-grouped training search; its encoder remains fixed. The original 256→64 head contains 279,943 trainable parameters. The author-recipe reference is logistic regression with C=71.68, LBFGS, maximum 1,000 iterations and seed 0, without a fitted scaler or class weighting.')
para('In the specialist branch, each model estimates a one-versus-rest score for one diagnosis. Sigmoid calibration and a binary threshold are fitted using out-of-fold training predictions. A validation-selected fusion strength adjusts the scores before the final argmax. The independent binary decision can identify zero or multiple classes, whereas the fused system selects one label; their sensitivities and F1 scores must therefore be reported separately.')
h('D. Software Components and Model Lifecycle',2)
para('The repository separates image loading and transforms (src/data), model definitions (src/models), training and metrics (src/utils), and the research dashboard (scripts/app.py and src/ui). Extraction caches contain embeddings and provenance; classifier folders store the fitted head, metadata, metric summaries, split-level predictions, and selection records. At inference, the selected backbone applies the associated preprocessing, and its saved classifier produces the seven class scores. The dashboard supports original and adapted checkpoints and shows aggregate and per-class results. Uploaded images are not screened by a validated lesion detector.')
h('E. Federated Branch',2)
para('Flower implements client-local updates, IID or Dirichlet label partitions, FedAvg, and optional FedProx. The default configuration specifies five clients and Dirichlet α=0.5. Opacus can add per-example gradient clipping, noise, and privacy accounting. Raw images remain local in the simulation. These are implementation options, not evidence of a deployed multi-hospital system or certified privacy. The saved demonstration uses two clients, two rounds, IID synthetic data, and differential privacy disabled.')
para('At each round, the server distributes current parameters, selected clients perform local optimization, and the server aggregates their model parameters using training-example weights. Conceptually, FedAvg computes w(t+1) = Σₖ(nₖ/N)wₖ(t+1) over participating clients, with N=Σₖnₖ. FedProx adds μ‖w−w(t)‖²/2 to each client objective to discourage excessive departure from the round’s starting model. The configuration enables the proximal term only when FedProx is selected. These equations describe the optional federated branch; the frozen-feature results below were produced centrally.')
para('When privacy is enabled, the client uses clipping and noise during gradient updates and tracks a cumulative image-level privacy budget across rounds. The README notes that validation/test statistics are separate evaluations and that the accounting does not establish a patient-level guarantee. A target ε in a configuration is not evidence of an achieved privacy budget for an unrun experiment; the saved synthetic demonstration has DP disabled.')
h('IV. EXPERIMENTAL PROTOCOL')
h('A. Evaluation Populations and Metrics',2)
para('The improved EfficientNet experiment documents 7,014 training, 1,509 internal validation, 1,492 internal test, 193 official validation, and 1,512 official test images. Training, validation, and internal test partitions are lesion-disjoint. The separate PanDerm released HAM split is 8,207 training, 575 validation, and 1,232 test images; these counts must not be combined with the EfficientNet split. External grouping identifiers do not support cross-release lesion or patient-independence claims.')
para('Accuracy is exact-label accuracy. Macro-F1 weights classes equally; weighted-F1 weights by class support; balanced accuracy averages class recall; macro-F0.5 emphasizes precision over recall. Macro metrics expose minority-class performance that accuracy can hide.')
para('For each class, precision is TP/(TP+FP), recall is TP/(TP+FN), and F1 is 2PR/(P+R), with zero-division handled consistently in the evaluation code. Macro-F1 is the arithmetic mean of the seven per-class F1 values; weighted-F1 weights them by the corresponding support. F0.5 uses (1+0.5²)PR/(0.5²P+R), giving precision more influence. A single accuracy number cannot indicate which diagnoses are missed or whether errors are concentrated in rare classes.')
h('B. Model Selection',2)
para('Improved EfficientNet selects representation and head settings using internal validation macro-F1, then freezes candidates before test scoring. PanDerm custom heads use either released validation or training-fold cross-validation as specified in their experiment. The authors’ probe is fixed. Since evaluation cohorts have been reused across the project history, these results are historical comparisons rather than prospective blind estimates.')
para('In the refined Large experiment, new candidates were compared using three stratified, lesion-disjoint folds inside the fitting partition. Each fold fitted its own scaler and, when applicable, its projection; the selected candidate was refitted with the median fold stopping budget. Out-of-fold scores used in selection are tuning diagnostics, not unbiased final performance estimates. Model files and prediction receipts were saved before each run’s test reporting, but the same cohorts had been inspected in earlier runs.')
h('V. RESULTS')
h('A. Improved EfficientNet Features',2)
para('Table I reports the improved frozen EfficientNet comparison. The linear classifier was selected by internal validation macro-F1 and achieved 0.6146 macro-F1 on the official test cohort. XGBoost had the strongest macro-F0.5 among these heads, so no family dominates every measure.')
table('TABLE I. IMPROVED EFFICIENTNET RESULTS (MACRO METRICS).',['Head','Val. F1','Internal F1','Official F1','Official F0.5'],[['Linear','0.6105','0.5976','0.6146','0.6150'],['MLP','0.6078','0.6083','0.6068','0.6105'],['XGBoost','0.5655','0.6020','0.5866','0.6171']])
para('On the same official cohort, the original EfficientNet linear model recorded 0.5385 macro-F1, while the improved linear model recorded 0.6146, a descriptive increase of 7.61 percentage points. The improved MLP reached 0.6068, compared with 0.5390 for its original version. These observations combine changes to views, tuning and related settings, so they do not isolate the effect of multi-view extraction alone. The prior experiment found that changing both SMOTE and class weighting prevented attribution of its XGBoost regression to SMOTE alone; a controlled validation ablation indicated that restoring weights recovered much of the lost score.')
h('B. PanDerm Comparisons',2)
para('Table II summarizes separate frozen PanDerm experiments on the official 1,512-image cohort. Base+MLP was validation-selected among Base heads. Large results used a different fitting split and selection protocol; they are not an architecture-only ablation. All figures are historical comparisons.')
table('TABLE II. PANDERM OFFICIAL-TEST METRICS (FRACTIONS).',['Encoder/head','Selection','Accuracy','Macro-F1','Macro-F0.5'],[['Base + MLP','Validation','0.8056','0.7030','0.7055'],['Base + linear','Validation','0.7917','0.6773','0.6865'],['Base + XGBoost','Validation','0.7566','0.6408','0.6558'],['Large + author probe','Fixed recipe','0.8294','0.7158','0.7415'],['Large + custom linear','Validation','0.8313','0.7380','0.7463'],['Large + MLP (192-64)','Train CV','0.8214','0.7096','0.7202'],['Large + XGBoost+LDA','Train CV','0.8201','0.7267','0.7235']])
para('On the distinct 1,232-image released HAM test split, the Large authors’ probe reached 92.37% accuracy, 80.98% macro-F1, and 92.22% weighted-F1. The Large 192-64 MLP reached 92.13% accuracy and 81.73% macro-F1, with 98.22% training macro-F1. The fit-to-test gap warrants caution; these results are not interchangeable with the official ISIC cohort.')
table('TABLE III. LARGE-ENCODER HEADS ON RELEASED HAM TEST (%).',['Head','Accuracy','Macro F1','Macro F0.5','Weighted F1'],[['Author probe','92.37','80.98','82.09','92.22'],['Custom linear','91.64','79.36','77.62','91.79'],['MLP 192–64','92.13','81.73','80.75','92.27'],['MLP 256–64','92.45','80.75','80.23','92.52'],['XGBoost+LDA','91.31','78.89','76.74','91.59']])
para('The five heads in Table III share the released Large encoder, but differ in classifier family and selection procedure. The 192–64 head leads this table on macro-F1; the author probe leads on macro-F0.5. This distinction matters because emphasizing precision changes the preferred model. The MLP 256–64 has the highest listed accuracy, yet its macro-F1 is below the 192–64 model’s. These descriptive differences do not establish statistically reliable superiority.')
h('C. Class-Level Performance',2)
para('Table IV shows the seven-class official-test performance for the frozen PanDerm Large custom linear head. Melanoma recall was 0.6901 and precision was 0.6277, while nevus F1 was 0.9129. Performance varies by lesion type; rare-class results are based on fewer cases and are more sensitive to a small number of predictions. These numbers are for the single seven-way prediction task and the historical 1,512-image cohort.')
table('TABLE IV. LARGE CUSTOM LINEAR: PER-CLASS OFFICIAL TEST.',['Class','Precision','Recall','F1','Support'],[['akiec','0.6410','0.5814','0.6098','43'],['bcc','0.7143','0.8065','0.7576','93'],['bkl','0.7561','0.7143','0.7346','217'],['df','0.7727','0.7727','0.7727','44'],['mel','0.6277','0.6901','0.6574','171'],['nv','0.9149','0.9109','0.9129','909'],['vasc','0.8462','0.6286','0.7213','35']])
h('D. Specialists and Regularization',2)
para('Regularized specialists achieved 0.5998 macro-F1 and 0.6453 macro-F0.5 on official test. The SMOTE counterpart reached 0.5858 and 0.6478. The improved EfficientNet linear model scored 0.6146 macro-F1 on the same repeatedly evaluated cohort. Specialization therefore did not improve overall macro-F1 over that reference. It represents a different per-class operating strategy, not a guaranteed overall gain.')
para('The older specialist system illustrates why task definitions matter: its independent melanoma specialist had 0.6667 recall on official test at its binary threshold, but the final seven-way combined decision had 0.4269 melanoma recall. Competition with other class scores changes the decision. The independent and combined figures should not be presented as two measurements of the same classifier output.')
h('E. Additional Head-Capacity Search',2)
para('A bounded PanDerm Large search compared residual corrections to a frozen linear control with plain neural heads. Across the new candidates, a single hidden layer of 128 units achieved 73.27% pooled tuning macro-F1 in training folds; a previously selected 192–64 head achieved 73.47% under the matched-fold comparison. The selected new 128-unit model scored 72.95% macro-F1 on official test versus 73.80% for the custom linear head. Its 95% paired bootstrap interval for the macro-F1 difference versus linear was −0.0371 to 0.0182. This interval is conditional on the chosen fitted models, excludes training and model-selection uncertainty, and includes zero; the search does not establish a unique optimal head size.')
h('F. Adaptation Scores',2)
para('The latest joint refinement applied eight further fitting updates to both former test cohorts (HAM n=1,232; official ISIC n=1,512). PanDerm Large 192-64 then scored 93.99% accuracy/87.38% macro-F1 on HAM and 85.91%/79.02% on ISIC. These are fitting scores, not test metrics. HAM validation was excluded from updates but had been inspected previously, so it is not fresh independent confirmation.')
para('The experiment gave each cohort 50% of total loss weight, used class-frequency power 0.25 and label smoothing 0.01, and attempted eight accumulated full-data updates for linear and neural models at learning rate 0.00014. The tree model appended boosting rounds at a different learning rate, so an update is not the same unit across families. A predeclared selection rule maximized the smaller improvement across the two cohorts in the arithmetic mean of accuracy, macro-F1, macro-F0.5, weighted-F1, and balanced accuracy. The selected 192–64 checkpoint was at additional update six; the rule itself uses the fitting cohorts and does not establish held-out superiority.')
table('TABLE V. JOINT-ADAPTATION COHORTS: FITTING METRICS (%).',['Large head','HAM acc.','HAM F1','ISIC acc.','ISIC F1'],[['192–64 MLP','93.99','87.38','85.91','79.02'],['128 MLP','93.26','85.75','85.71','78.60'],['Linear','92.53','81.74','84.33','76.19'],['XGBoost+LDA','91.48','79.18','82.21','73.08'],['Author probe*','92.37','80.98','82.94','71.58']])
para('*The author-probe reference was unchanged; the four custom heads were selected after fitting to the cohorts. Scores are descriptive of those same cohorts. Three Large heads exceeded the reference on the specified five-metric mean on both cohorts; that composite does not imply that every individual metric improved.')
h('G. Federated Demonstration',2)
para('The saved two-client, two-round synthetic FedAvg run used one local epoch and no DP. Final accuracy was 14.29%, nominal chance for seven balanced categories. This verifies pipeline execution only. No real-image FL performance, privacy-utility result, or matched centralized comparison is established by this receipt.')
h('VI. DISCUSSION')
para('Frozen pretrained representations support useful lightweight classifiers, but accuracy and weighted-F1 alone obscure minority-diagnosis errors. PanDerm results vary with backbone scale, probe design, fitting cohort, and selection procedure. The Large custom linear result is a promising historical comparison but does not isolate backbone scale causally from the Base result. The local authors’ probe follows released implementation artifacts; it is not a claim of exact manuscript reproduction.')
para('The strongest centralized scores did not come from federated or differentially private training. A real FL study needs institution-level partitions, client-specific test results, measured privacy budgets, and a matched centralized baseline.')
para('A comparison of large versus base backbones should also hold the fitting split, training budget, head family, and selection rule constant. Current Large and Base comparisons differ on several of these factors. The local author reference likewise follows the released CSV and code, whose split proportions differ from those described in the manuscript; the reproduction is a local recipe reference rather than a bitwise recreation of the publication. The paper’s weighted-F1 and local macro-F1 describe different summaries and cannot be substituted for each other.')
para('Uncertainty estimates are sensitive to the unit of resampling. One internal-test comparison in the project used lesion-level bootstrap resampling to account for related images; the official test files do not provide lesion groups, so the corresponding capacity-search interval resampled individual rows. Both fixed-model intervals omit the variability introduced by retraining and selecting models. Small differences should be interpreted alongside uncertainty, sample support, and the history of repeated cohort inspection.')
h('VII. LIMITATIONS AND RESPONSIBLE USE')
para('The prototype has no prospective clinical validation, multi-institution evaluation, fairness assessment, or workflow study. Pretraining overlap has not been independently excluded. Repeated inspection of evaluation cohorts creates adaptive-reuse risk even if a particular run did not use test labels for selection. Cross-release patient/lesion independence is unknown. Differential privacy was disabled in the saved federated run; image-level accounting would not by itself prove patient-level protection. Decision scores are not clinical probabilities. The upload checker does not verify that an input depicts a lesion. This is not a medical device or diagnostic tool.')
h('VIII. CONCLUSION')
para('The project provides an auditable framework for comparing frozen visual encoders with lightweight classifiers for seven-class lesion recognition. The strongest historical official-cohort macro-F1 is 0.7380 from a validation-selected PanDerm Large linear head; PanDerm Base+MLP reached 0.7030 and improved EfficientNet linear features reached 0.6146. These are protocol-dependent, repeatedly inspected measurements, not fresh clinical evidence. Adapted-cohort scores are fitting diagnostics, and the federated system has only been smoke-tested with synthetic data. Priorities are an untouched external cohort, matched encoder/head protocols, patient-level grouping, uncertainty reporting, and real multi-institution federated evaluation.')
h('IX. REPRODUCIBILITY AND ARTIFACTS')
para('Experiment reports and receipts are stored in the project’s feature-model output folders for improved EfficientNet, PanDerm Base, PanDerm Large, refined heads and the capacity search. The joint-refinement diagnostics folder records the adapted-head protocol and selection receipts. These folders include split manifests, configuration and selection metadata, saved predictions, per-class statistics, and verification receipts. The federated-demo artifacts contain the synthetic run’s metrics.')
para('The recorded project commands are scripts/train_improved.py and scripts/report_improved.py for the improved EfficientNet pipeline; scripts/train_panderm_large_heads.py and scripts/report_panderm_large.py for the initial Large comparison; and train-panderm-refined.cmd and train-panderm-capacity.cmd for the subsequent head studies. These workflows require the corresponding local data, checkpoints, dependencies, and hardware. Reproduction should preserve the declared split and model-selection procedure, and any new confirmation should use a cohort that has not been used for adaptation or repeated selection.')
h('ACKNOWLEDGMENT')
para('Prepared from repository source, experiment reports, metric receipts, and verification artifacts available October 28, 2026.')
h('REFERENCES')
refs=[
'[1] P. Tschandl, C. Rosendahl, and H. Kittler, “The HAM10000 dataset, a large collection of multi-source dermatoscopic images of common pigmented skin lesions,” Scientific Data, vol. 5, Art. no. 180161, 2018, doi: 10.1038/sdata.2018.161.',
'[2] N. Codella et al., “Skin lesion analysis toward melanoma detection 2018: A challenge hosted by the International Skin Imaging Collaboration,” arXiv:1902.03368, 2019. Available: https://arxiv.org/abs/1902.03368',
'[3] S. Yan et al., “A multimodal vision foundation model for clinical dermatology,” Nature Medicine, vol. 31, pp. 2691–2702, 2025, doi: 10.1038/s41591-025-03747-y.',
'[4] M. Tan and Q. V. Le, “EfficientNet: Rethinking model scaling for convolutional neural networks,” in Proc. 36th Int. Conf. Machine Learning, 2019, pp. 6105–6114.',
'[5] AI Project Research Team, “Federated Privacy-Preserving Skin Lesion Classifier: source code and experiment receipts,” project repository, accessed Oct. 28, 2026.'
]
for x in refs:
    p=doc.add_paragraph(x); p.paragraph_format.left_indent=Inches(.14); p.paragraph_format.first_line_indent=Inches(-.14); p.paragraph_format.space_after=Pt(2)
    for r in p.runs: r.font.size=Pt(8)
# A trailing continuous section break tells Word to balance the final two-column page.
endsec=doc.add_section(WD_SECTION.CONTINUOUS)
endcols=endsec._sectPr.find(qn('w:cols'))
if endcols is None: endcols=OxmlElement('w:cols'); endsec._sectPr.append(endcols)
endcols.set(qn('w:num'),'1')
f=doc.sections[0].footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
f.add_run('IEEE-style project report  |  ').font.size=Pt(8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
doc.core_properties.title='A Comparative Study of Frozen Visual Representations for Seven-Class Dermoscopic Lesion Classification'
doc.core_properties.author='AI Project Research Team'
doc.save(out); print(out.resolve())
